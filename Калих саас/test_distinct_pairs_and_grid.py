"""
test_distinct_pairs_and_grid.py — Комплексная валидация новой сетки звонков и правила запрета сдваивания пар.

Проверяет:
a) Каждый день состоит из ровно 3 пар и 1 одиночного урока (слоты 1, 2, 3, 4).
b) В любой учебный день слоты 1, 2 и 3 (пары) содержат строго 3 РАЗНЫХ предмета (subject_id).
c) Экспорт в Excel (Sheet 1) содержит ровно 4 строки звонков (3 пары + 1 урок),
   и 4-я строка (4-й урок) содержит распределенные занятия.
d) Ни один предмет не превышает 3 академических часа в один день.
"""

import unittest
from io import BytesIO
from datetime import date
from collections import defaultdict
import openpyxl

from app import app
from models import db, Lyceum, StudentGroup, ScheduleAssignment, SlotType
from excel_exporter import generate_excel_matrix
from seed_db import seed_database


class TestDistinctPairsAndGrid(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with app.app_context():
            seed_database()

    def setUp(self):
        self.app_context = app.app_context()
        self.app_context.push()

    def tearDown(self):
        self.app_context.pop()

    def test_daily_bells_grid_structure(self):
        """a) Проверка сетки звонков: каждый день содержит слоты 1, 2, 3 (пары) и 4 (урок)."""
        assignments = ScheduleAssignment.query.order_by(ScheduleAssignment.date, ScheduleAssignment.slot_number).all()
        self.assertGreater(len(assignments), 0, "Расписание не должно быть пустым")

        days_map = defaultdict(dict)
        for a in assignments:
            days_map[a.date][a.slot_number] = a

        # Проверяем структуру каждого активного дня
        full_days_count = 0
        partial_days = []
        for current_date, slots in days_map.items():
            if len(slots) == 4:
                full_days_count += 1
                self.assertEqual(set(slots.keys()), {1, 2, 3, 4}, f"В полный день {current_date} должны быть слоты 1, 2, 3, 4")
            else:
                partial_days.append((current_date, slots))

            # Временные интервалы согласно сетке звонков
            if 1 in slots:
                self.assertEqual(slots[1].slot_type, SlotType.PAIR)
                self.assertEqual(slots[1].start_time.strftime("%H:%M"), "08:30")
                self.assertEqual(slots[1].end_time.strftime("%H:%M"), "10:00")
            if 2 in slots:
                self.assertEqual(slots[2].slot_type, SlotType.PAIR)
                self.assertEqual(slots[2].start_time.strftime("%H:%M"), "10:10")
                self.assertEqual(slots[2].end_time.strftime("%H:%M"), "11:40")
            if 3 in slots:
                self.assertEqual(slots[3].slot_type, SlotType.PAIR)
                self.assertEqual(slots[3].start_time.strftime("%H:%M"), "12:00")
                self.assertEqual(slots[3].end_time.strftime("%H:%M"), "13:30")
            if 4 in slots:
                self.assertEqual(slots[4].slot_type, SlotType.SINGLE_LESSON)
                self.assertEqual(slots[4].start_time.strftime("%H:%M"), "13:40")
                self.assertEqual(slots[4].end_time.strftime("%H:%M"), "14:25")

        self.assertEqual(len(days_map), 206, "Всего учебных дней с занятиями должно быть ровно 206")
        self.assertEqual(full_days_count, 205, "Ровно 205 дней должны иметь полную сетку из 4 слотов (7 акад. часов)")
        self.assertEqual(len(partial_days), 1, "Только 1 финальный день может быть неполным из-за 1440 mod 7 = 5ч")
        final_date, final_slots = partial_days[0]
        self.assertEqual(set(final_slots.keys()), {1, 2, 4}, f"Финальный день {final_date} содержит слоты 1, 2, 4 (5 часов)")

    def test_anti_doubling_distinct_pairs_rule(self):
        """b) Запрет сдваивания пар: в любой день слоты 1, 2, 3 имеют строго РАЗНЫЕ subject_id."""
        assignments = ScheduleAssignment.query.all()
        days_map = defaultdict(dict)
        for a in assignments:
            days_map[a.date][a.slot_number] = a

        for current_date, slots in days_map.items():
            pair_slots = [slots[s] for s in [1, 2, 3] if s in slots]
            subject_ids = [p.subject_id for p in pair_slots]
            subject_names = [p.subject.name for p in pair_slots if p.subject]

            # Проверка уникальности предметов среди пар
            self.assertEqual(
                len(subject_ids), len(set(subject_ids)),
                f"ОБНАРУЖЕНО СДВАИВАНИЕ ПАР в день {current_date}: {subject_names}"
            )
            if len(pair_slots) == 3:
                self.assertEqual(len(set(subject_ids)), 3, f"3 пары должны быть 3 РАЗНЫМИ предметами на {current_date}")

            # Проверка лимита: максимум 3 академических часа на предмет в день
            daily_hours = defaultdict(int)
            for s_num, a in slots.items():
                hrs = 2 if a.slot_type == SlotType.PAIR else 1
                daily_hours[a.subject_id] += hrs

            for subj_id, h in daily_hours.items():
                self.assertLessEqual(
                    h, 3,
                    f"Предмет ID {subj_id} превысил дневной лимит в 3ч на дату {current_date}: факт {h}ч"
                )

    def test_excel_export_sheet1_and_sheet2_integrity(self):
        """c) Экспорт Excel: Sheet 1 имеет 4 строки звонков с занятиями в 4 строке, Sheet 2 содержит все слоты."""
        lyceum = Lyceum.query.first()
        self.assertIsNotNone(lyceum)

        buf = generate_excel_matrix(lyceum.id)
        wb = openpyxl.load_workbook(buf)

        sheet1_name = "Группа ВЕБ-25"
        self.assertIn(sheet1_name, wb.sheetnames)
        ws1 = wb[sheet1_name]

        # 4 строки сетки звонков:
        # Row 5: 1-я пара (08:30 - 10:00)
        # Row 6: 2-я пара (10:10 - 11:40)
        # Row 7: 3-я пара (12:00 - 13:30)
        # Row 8: 4-й урок (13:40 - 14:25)
        self.assertIn("1-я пара", str(ws1["A5"].value))
        self.assertIn("08:30 - 10:00", str(ws1["A5"].value))

        self.assertIn("2-я пара", str(ws1["A6"].value))
        self.assertIn("10:10 - 11:40", str(ws1["A6"].value))

        self.assertIn("3-я пара", str(ws1["A7"].value))
        self.assertIn("12:00 - 13:30", str(ws1["A7"].value))

        self.assertIn("4-й урок", str(ws1["A8"].value))
        self.assertIn("13:40 - 14:25", str(ws1["A8"].value))

        # Проверяем, что в Sheet 1 для каждого дня недели (Пн-Сб) все 3 пары (строки 5, 6, 7) имеют РАЗНЫЕ предметы
        for col_letter in ["B", "C", "D", "E", "F", "G"]:
            pair_subjects = []
            for r in [5, 6, 7]:
                val = str(ws1[f"{col_letter}{r}"].value)
                self.assertNotEqual(val, "—", f"Ячейка {col_letter}{r} должна быть заполнена")
                subj_name = val.split("\n")[0]
                pair_subjects.append(subj_name)
            self.assertEqual(
                len(pair_subjects), len(set(pair_subjects)),
                f"В Excel Sheet 1 для колонки {col_letter} пары должны иметь разные предметы: {pair_subjects}"
            )

        # Проверяем, что 4-я строка (4-й урок) содержит реальные занятия по всем дням Пн-Сб (B8:G8)
        for col_letter in ["B", "C", "D", "E", "F", "G"]:
            cell_val = str(ws1[f"{col_letter}8"].value)
            self.assertNotEqual(cell_val, "—", f"Ячейка {col_letter}8 (4-й урок) должна быть заполнена")
            self.assertIn("\n", cell_val, f"Ячейка {col_letter}8 должна содержать многострочную информацию")

        # Проверка Sheet 2 (Годовой журнал)
        sheet2_name = "Журнал ВЕБ-25 (Год)"
        self.assertIn(sheet2_name, wb.sheetnames)
        ws2 = wb[sheet2_name]

        assignments = ScheduleAssignment.query.filter_by(lyceum_id=lyceum.id).all()
        # В Sheet 2 число строк с занятиями равно общему числу assignments (823)
        populated_count = 0
        single_lesson_count = 0
        pair_count = 0

        for r in range(5, 5 + len(assignments)):
            slot_label = str(ws2[f"D{r}"].value)
            if "урок" in slot_label:
                single_lesson_count += 1
            elif "пара" in slot_label:
                pair_count += 1
            populated_count += 1

        self.assertEqual(populated_count, len(assignments))
        self.assertGreater(single_lesson_count, 0, "В Sheet 2 должны присутствовать одиночные уроки")
        self.assertGreater(pair_count, 0, "В Sheet 2 должны присутствовать пары")

        # Итоговая сумма часов в последней строке Sheet 2
        sum_row = 5 + len(assignments)
        total_hours_in_sheet = int(ws2[f"I{sum_row}"].value)
        self.assertEqual(total_hours_in_sheet, 1440, "Итого часов в Sheet 2 должно быть ровно 1440")


if __name__ == "__main__":
    unittest.main()
