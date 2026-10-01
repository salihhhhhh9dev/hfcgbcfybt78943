"""
Script to create Excel file for Official Ministry Curriculum Plan (1st Year, 1440 Hours)
Profession: «Разработчик web и мультимедийных приложений» (2513 Разработчик web-сайта, 3121 Web дизайнер).
Group: «ВЕБ-25» (1 курс, 2 года обучения).
Tech Stack: Python, openpyxl
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


CURRICULUM_DATA = [
    # (Category, Subject Name, Hours, Teacher, Room, Form of Control)
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Кыргызский язык", 102, "Асанова Айгуль Бакытовна", "Каб. 201", "Экзамен"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Кыргызская литература", 34, "Асанова Айгуль Бакытовна", "Каб. 201", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Русский язык", 68, "Смирнова Елена Васильевна", "Каб. 202", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Русская литература", 102, "Смирнова Елена Васильевна", "Каб. 202", "Экзамен"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Иностранный язык (Английский)", 68, "Айтматова Назгуль Чынгызовна", "Каб. 201", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "История", 68, "Жумабеков Канат Темирович", "Каб. 203", "Экзамен"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Человек и общество", 34, "Жумабеков Канат Темирович", "Каб. 203", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Алгебра", 102, "Исаков Данияр Алымкулович", "Каб. 201", "Экзамен"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Геометрия", 34, "Исаков Данияр Алымкулович", "Каб. 201", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Биология", 34, "Токтосунова Жылдыз Муратовна", "Каб. 202", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "География", 34, "Токтосунова Жылдыз Муратовна", "Каб. 202", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Физика / Астрономия", 102, "Токтосунова Жылдыз Муратовна", "Каб. 202", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Химия", 68, "Токтосунова Жылдыз Муратовна", "Каб. 202", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Допризывная подготовка", 68, "Мамбетов Кубат Садыкович", "Каб. 203", "Зачет"),
    ("I. ОБЩЕОБРАЗОВАТЕЛЬНЫЙ КОМПОНЕНТ", "Физическое воспитание", 34, "Мамбетов Кубат Садыкович", "Спортзал", "Зачет"),

    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Охрана труда и производственная экология", 24, "Бакиров Руслан Эркинович", "IT-мастерская (веб-разработка)", "Зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Правовые основы профессиональной деятельности", 24, "Жумабеков Канат Темирович", "Каб. 203", "Зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Безопасный здоровый образ жизни", 24, "Мамбетов Кубат Садыкович", "Каб. 203", "Зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Основы компьютерных вычислений", 42, "Исаков Данияр Алымкулович", "Компьютерная лаборатория №1", "Зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Библиотека языков программирования", 30, "Калиев Самат Бектурович", "Компьютерная лаборатория №1", "Дифф. зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Основы программирования", 60, "Калиев Самат Бектурович", "Компьютерная лаборатория №1", "Экзамен"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Технология программирования", 40, "Калиев Самат Бектурович", "Компьютерная лаборатория №1", "Дифф. зачет"),
    ("II. ПРОФЕССИОНАЛЬНЫЙ КОМПОНЕНТ", "Производственное обучение (Практикум в лаборатории)", 244, "Бакиров Руслан Эркинович", "IT-мастерская (веб-разработка)", "Экзамен"),
]


def create_excel_curriculum():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Учебный план ВЕБ-25 (1440ч)"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    font_title = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Arial", size=10, italic=True, color="475569")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_category = Font(name="Arial", size=10, bold=True, color="0F172A")
    font_data = Font(name="Arial", size=9, color="0F172A")
    font_summary = Font(name="Arial", size=10, bold=True, color="0F172A")

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_category = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_summary = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # Title Banner (Rows 1-3)
    ws.merge_cells("A1:F1")
    ws.cell(row=1, column=1, value="МИНИСТЕРСТВО ПРОСВЕЩЕНИЯ КЫРГЫЗСКОЙ РЕСПУБЛИКИ").font = font_title
    ws.cell(row=1, column=1).fill = fill_title
    ws.cell(row=1, column=1).alignment = align_center
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws.cell(row=2, column=1, value="ТИПОВОЙ УЧЕБНЫЙ ПЛАН — 1 КУРС (Группа ВЕБ-25, 2026/2027 уч. год)").font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws.cell(row=2, column=1).fill = fill_title
    ws.cell(row=2, column=1).alignment = align_center
    ws.row_dimensions[2].height = 24

    ws.merge_cells("A3:F3")
    ws.cell(row=3, column=1, value="Профессия: Разработчик web и мультимедийных приложений | Квалификации: 2513 Web-разработчик, 3121 Web-дизайнер | 36 ч/нед").font = font_subtitle
    ws.cell(row=3, column=1).alignment = align_center
    ws.row_dimensions[3].height = 20

    # Headers (Row 5)
    ws.row_dimensions[4].height = 8
    ws.row_dimensions[5].height = 26
    headers = ["№", "Наименование учебной дисциплины", "Всего часов", "Преподаватель", "Кабинет / Лаборатория", "Форма контроля"]
    widths = [6, 48, 14, 30, 32, 16]

    for col_idx, (h_text, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=5, column=col_idx, value=h_text)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = cell_border
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    current_row = 6
    current_cat = None
    item_num = 1
    total_hours_sum = 0
    cat_hours_sum = 0

    for cat, name, hours, teacher, room, ctrl in CURRICULUM_DATA:
        if cat != current_cat:
            if current_cat is not None:
                # Subtotal for previous category
                ws.row_dimensions[current_row].height = 22
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
                cs1 = ws.cell(row=current_row, column=1, value=f"ИТОГО ПО РАЗДЕЛУ ({current_cat}):")
                cs2 = ws.cell(row=current_row, column=3, value=cat_hours_sum)
                for col in range(1, 7):
                    ws.cell(row=current_row, column=col).border = cell_border
                    ws.cell(row=current_row, column=col).fill = fill_category
                cs1.font = font_category; cs1.alignment = align_left
                cs2.font = font_category; cs2.alignment = align_center
                current_row += 1

            current_cat = cat
            cat_hours_sum = 0
            # Category Header Row
            ws.row_dimensions[current_row].height = 24
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
            c_cat = ws.cell(row=current_row, column=1, value=cat)
            c_cat.font = font_category
            c_cat.fill = fill_category
            c_cat.alignment = align_left
            for col in range(1, 7):
                ws.cell(row=current_row, column=col).border = cell_border
                ws.cell(row=current_row, column=col).fill = fill_category
            current_row += 1

        cat_hours_sum += hours
        total_hours_sum += hours

        ws.row_dimensions[current_row].height = 20
        c_num = ws.cell(row=current_row, column=1, value=item_num)
        c_name = ws.cell(row=current_row, column=2, value=name)
        c_hours = ws.cell(row=current_row, column=3, value=hours)
        c_t = ws.cell(row=current_row, column=4, value=teacher)
        c_r = ws.cell(row=current_row, column=5, value=room)
        c_ctrl = ws.cell(row=current_row, column=6, value=ctrl)

        for col, c_cell in enumerate([c_num, c_name, c_hours, c_t, c_r, c_ctrl], start=1):
            c_cell.font = font_data
            c_cell.border = cell_border
            if item_num % 2 == 0:
                c_cell.fill = fill_alt

        c_num.alignment = align_center
        c_name.alignment = align_left
        c_hours.alignment = align_center
        c_t.alignment = align_left
        c_r.alignment = align_left
        c_ctrl.alignment = align_center

        item_num += 1
        current_row += 1

    # Subtotal for last category
    ws.row_dimensions[current_row].height = 22
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    cs1 = ws.cell(row=current_row, column=1, value=f"ИТОГО ПО РАЗДЕЛУ ({current_cat}):")
    cs2 = ws.cell(row=current_row, column=3, value=cat_hours_sum)
    for col in range(1, 7):
        ws.cell(row=current_row, column=col).border = cell_border
        ws.cell(row=current_row, column=col).fill = fill_category
    cs1.font = font_category; cs1.alignment = align_left
    cs2.font = font_category; cs2.alignment = align_center
    current_row += 1

    # Grand Total Row
    ws.row_dimensions[current_row].height = 26
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    gt1 = ws.cell(row=current_row, column=1, value="ВСЕГО НАГРУЗКА 1-ГО КУРСА (36 ЧАСОВ В НЕДЕЛЮ):")
    gt2 = ws.cell(row=current_row, column=3, value=total_hours_sum)
    for col in range(1, 7):
        ws.cell(row=current_row, column=col).border = cell_border
        ws.cell(row=current_row, column=col).fill = fill_summary
    gt1.font = font_summary; gt1.alignment = align_left
    gt2.font = font_summary; gt2.alignment = align_center

    target_path = "Учебный_План_1440_часов_ВЕБ-25.xlsx"
    wb.save(target_path)
    print(f"Created '{target_path}' successfully ({len(CURRICULUM_DATA)} subjects, Total: {total_hours_sum} hours).")


if __name__ == "__main__":
    create_excel_curriculum()
