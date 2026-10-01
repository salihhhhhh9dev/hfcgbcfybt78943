"""
Excel Exporter Utility for Vocational Lyceum B2B SaaS
Tech Stack: openpyxl, Flask, SQLAlchemy
"""

import io
from datetime import datetime, date, time, timedelta
from typing import Dict, List
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import db, ScheduleAssignment, Lyceum, StudentGroup, Teacher, Room, SubjectCurriculum, LyceumSettings, SlotType

# Day Mapping (Columns B-G)
DAYS_MAP = [
    (1, "Понедельник"),
    (2, "Вторник"),
    (3, "Среда"),
    (4, "Четверг"),
    (5, "Пятница"),
    (6, "Суббота")
]


def generate_excel_matrix(lyceum_id: int) -> io.BytesIO:
    """
    Generates a professionally formatted Excel timetable matrix.
    Creates a separate Worksheet for each Group.
    Columns = Days of the Week (Пн-Сб).
    Rows = Time Slots (Pairs & Single Lessons).
    Intersecting cells = "Subject Name \\n Teacher Name \\n Room".
    Returns in-memory io.BytesIO buffer.
    """
    # 1. Fetch Schedule Data with relations
    slots = (
        ScheduleAssignment.query
        .filter(ScheduleAssignment.lyceum_id == lyceum_id)
        .order_by(ScheduleAssignment.date.asc(), ScheduleAssignment.slot_number.asc())
        .all()
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # 2. Group slots by Group entity
    grouped_slots: Dict[StudentGroup, List[ScheduleAssignment]] = defaultdict(list)
    for slot in slots:
        if slot.group:
            grouped_slots[slot.group].append(slot)

    # If no slots exist yet, fetch all groups to create empty templates
    if not grouped_slots:
        all_groups = StudentGroup.query.filter(StudentGroup.lyceum_id == lyceum_id).all()
        for g in all_groups:
            grouped_slots[g] = []

    # If still no groups found, create default summary tab
    if not grouped_slots:
        ws = wb.create_sheet(title="Инфо")
        ws.cell(row=1, column=1, value="Расписание не найдено для данного лицея.").font = Font(bold=True)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # Determine dynamic time slots from LyceumSettings
    settings = LyceumSettings.query.filter_by(lyceum_id=lyceum_id).first()
    pairs_cnt = settings.pairs_per_day if (settings and settings.pairs_per_day) else 3
    single_cnt = settings.single_lessons_per_day if (settings and settings.single_lessons_per_day is not None and settings.single_lessons_per_day > 0) else 1
    start_t = settings.day_start_time if settings and settings.day_start_time else time(8, 30)
    p_dur = settings.pair_duration_minutes if settings and settings.pair_duration_minutes else 90
    s_dur = settings.single_lesson_duration_minutes if settings and settings.single_lesson_duration_minutes else 45
    reg_break = settings.regular_break_minutes if settings and settings.regular_break_minutes else 10
    long_break_slot = settings.long_break_after_pair if settings and settings.long_break_after_pair else 2
    long_break = settings.long_break_minutes if settings and settings.long_break_minutes else 20

    dynamic_time_slots = []
    cur_dt = datetime.combine(date(2026, 1, 1), start_t)
    s_num = 1
    for _ in range(pairs_cnt):
        end_dt = cur_dt + timedelta(minutes=p_dur)
        dynamic_time_slots.append({
            "slot_number": s_num,
            "slot_label": f"{s_num}-я пара",
            "time_str": f"{cur_dt.strftime('%H:%M')} - {end_dt.strftime('%H:%M')}"
        })
        pause = long_break if s_num == long_break_slot else reg_break
        cur_dt = end_dt + timedelta(minutes=pause)
        s_num += 1

    for _ in range(single_cnt):
        end_dt = cur_dt + timedelta(minutes=s_dur)
        dynamic_time_slots.append({
            "slot_number": s_num,
            "slot_label": f"{s_num}-й урок",
            "time_str": f"{cur_dt.strftime('%H:%M')} - {end_dt.strftime('%H:%M')}"
        })
        cur_dt = end_dt + timedelta(minutes=reg_break)
        s_num += 1

    if not dynamic_time_slots:
        dynamic_time_slots = [
            {"slot_number": 1, "slot_label": "1-я пара", "time_str": "08:30 - 10:00"},
            {"slot_number": 2, "slot_label": "2-я пара", "time_str": "10:10 - 11:40"},
            {"slot_number": 3, "slot_label": "3-я пара", "time_str": "12:00 - 13:30"},
            {"slot_number": 4, "slot_label": "4-й урок", "time_str": "13:40 - 14:25"}
        ]

    # Reusable OpenPyXL Styling
    font_title = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=10, italic=True, color="64748B")
    font_header = Font(name="Arial", size=11, bold=True, color="1E293B")
    font_time = Font(name="Arial", size=10, bold=True, color="334155")
    font_content = Font(name="Arial", size=10, color="0F172A")

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_header = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_time = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_slot = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    border_thin = Side(border_style="thin", color="CBD5E1")
    border_thick = Side(border_style="medium", color="475569")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    align_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")

    # 3. Build Timetable Sheet per Group
    for group, group_slots in grouped_slots.items():
        sheet_title = f"Группа {group.name}"[:31]
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Header Title Banner (Merged A1:G1)
        ws.merge_cells("A1:G1")
        title_cell = ws.cell(row=1, column=1, value=f"РАСПИСАНИЕ ЗАНЯТИЙ — ГРУППА {group.name} ({group.course_year} курс)")
        title_cell.font = font_title
        title_cell.fill = fill_title
        title_cell.alignment = align_center
        ws.row_dimensions[1].height = 36

        # Subtitle info (Merged A2:G2)
        ws.merge_cells("A2:G2")
        sub_cell = ws.cell(row=2, column=1, value=f"Студентов в группе: {group.student_count} | Всего занятий в базе: {len(group_slots)}")
        sub_cell.font = font_sub
        sub_cell.alignment = align_center
        ws.row_dimensions[2].height = 20

        # Row 3: Blank separator
        ws.row_dimensions[3].height = 10

        # Row 4: Column Headers (A4 = "Время / Пара", B4:G4 = Days of Week)
        ws.row_dimensions[4].height = 28
        c_corner = ws.cell(row=4, column=1, value="Время / Пара")
        c_corner.font = font_header
        c_corner.fill = fill_header
        c_corner.alignment = align_center
        c_corner.border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thick)

        for idx, (day_code, day_name) in enumerate(DAYS_MAP, start=2):
            c_day = ws.cell(row=4, column=idx, value=day_name.upper())
            c_day.font = font_header
            c_day.fill = fill_header
            c_day.alignment = align_center
            c_day.border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thick)

        # Index group slots by (day_of_week_int, slot_number)
        slot_matrix = {}
        for s in group_slots:
            d_val = int(s.day_of_week.value if hasattr(s.day_of_week, "value") else s.day_of_week)
            # Use first occurrence or representative slot
            if (d_val, s.slot_number) not in slot_matrix:
                slot_matrix[(d_val, s.slot_number)] = s

        # Build Timetable Rows using dynamic time slots
        for r_idx, slot_info in enumerate(dynamic_time_slots, start=5):
            s_num = slot_info["slot_number"]
            slot_label = slot_info["slot_label"]
            time_str = slot_info["time_str"]

            ws.row_dimensions[r_idx].height = 58  # Height for 3 lines of wrapped text

            # Column A: Time Slot Label
            c_time = ws.cell(row=r_idx, column=1, value=f"{slot_label}\n{time_str}")
            c_time.font = font_time
            c_time.fill = fill_time
            c_time.alignment = align_center_wrap
            c_time.border = cell_border

            # Columns B-G: Intersecting Cells per Day
            for c_idx, (day_code, _) in enumerate(DAYS_MAP, start=2):
                matched_slot = slot_matrix.get((day_code, s_num))

                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = cell_border
                cell.alignment = align_center_wrap
                cell.font = font_content

                if matched_slot:
                    t_name = matched_slot.teacher.full_name if matched_slot.teacher else "—"
                    r_name = matched_slot.room.room_number if matched_slot.room else "—"
                    s_name = matched_slot.subject.name if matched_slot.subject else "Занятие"
                    cell_text = f"{s_name}\n{t_name}\n{r_name}"
                    cell.value = cell_text
                    cell.fill = fill_slot
                else:
                    cell.value = "—"

        # 4. Auto-fit Column Widths cleanly
        ws.column_dimensions["A"].width = 18
        end_row = 5 + len(dynamic_time_slots)
        for col_idx in range(2, 8):
            col_letter = get_column_letter(col_idx)
            max_len = 16  # minimum column width for day columns
            for r in range(5, end_row):
                val = ws.cell(row=r, column=col_idx).value
                if val and val != "—":
                    lines = str(val).split("\n")
                    max_len = max(max_len, max(len(line) for line in lines))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 22)

        # 4b. Also build Annual Register Worksheet with all individual scheduled lessons
        if group_slots:
            ws_annual = wb.create_sheet(title=f"Журнал {group.name} (Год)"[:31])
            ws_annual.views.sheetView[0].showGridLines = True

            # Banner
            ws_annual.merge_cells("A1:I1")
            ann_title = ws_annual.cell(row=1, column=1, value=f"ГОДОВОЙ РЕЕСТР ЗАНЯТИЙ — ГРУППА {group.name} ({group.course_year} курс)")
            ann_title.font = font_title
            ann_title.fill = fill_title
            ann_title.alignment = align_center
            ws_annual.row_dimensions[1].height = 32

            ws_annual.merge_cells("A2:I2")
            total_slot_hours = sum(2 if (s.slot_type == SlotType.PAIR or getattr(s.slot_type, "value", s.slot_type) == "pair") else 1 for s in group_slots)
            ann_sub = ws_annual.cell(row=2, column=1, value=f"Всего занятий в плане: {len(group_slots)} | Суммарная нагрузка: {total_slot_hours} академических часов")
            ann_sub.font = font_sub
            ann_sub.alignment = align_center
            ws_annual.row_dimensions[2].height = 20

            # Headers
            ws_annual.row_dimensions[4].height = 26
            ann_headers = [
                ("A", "№", 8),
                ("B", "Дата", 14),
                ("C", "День недели", 18),
                ("D", "Пара / Урок", 16),
                ("E", "Время", 16),
                ("F", "Учебная дисциплина", 42),
                ("G", "Преподаватель", 32),
                ("H", "Аудитория / Лаборатория", 32),
                ("I", "Часов", 10),
            ]
            for col_l, h_txt, col_w in ann_headers:
                hc = ws_annual[f"{col_l}4"]
                hc.value = h_txt
                hc.font = font_header
                hc.fill = fill_header
                hc.alignment = align_center
                hc.border = cell_border
                ws_annual.column_dimensions[col_l].width = col_w

            day_names_ru = {
                1: "Понедельник", 2: "Вторник", 3: "Среда",
                4: "Четверг", 5: "Пятница", 6: "Суббота", 7: "Воскресенье"
            }

            for row_i, slot in enumerate(group_slots, start=5):
                ws_annual.row_dimensions[row_i].height = 20
                d_int = int(slot.day_of_week.value if hasattr(slot.day_of_week, "value") else slot.day_of_week)
                d_name = day_names_ru.get(d_int, str(d_int))
                is_pair = (slot.slot_type == SlotType.PAIR or getattr(slot.slot_type, "value", slot.slot_type) == "pair")
                slot_hrs = 2 if is_pair else 1
                type_label = f"{slot.slot_number}-я пара" if is_pair else f"{slot.slot_number}-й урок"
                time_range = f"{slot.start_time.strftime('%H:%M')} - {slot.end_time.strftime('%H:%M')}"
                subj_name = slot.subject.name if slot.subject else "—"
                teach_name = slot.teacher.full_name if slot.teacher else "—"
                room_name = slot.room.room_number if slot.room else "—"

                c_num = ws_annual.cell(row=row_i, column=1, value=row_i - 4)
                c_date = ws_annual.cell(row=row_i, column=2, value=slot.date.strftime("%d.%m.%Y"))
                c_day = ws_annual.cell(row=row_i, column=3, value=d_name)
                c_slot = ws_annual.cell(row=row_i, column=4, value=type_label)
                c_time = ws_annual.cell(row=row_i, column=5, value=time_range)
                c_subj = ws_annual.cell(row=row_i, column=6, value=subj_name)
                c_teach = ws_annual.cell(row=row_i, column=7, value=teach_name)
                c_room = ws_annual.cell(row=row_i, column=8, value=room_name)
                c_hrs = ws_annual.cell(row=row_i, column=9, value=slot_hrs)

                fill_curr = fill_slot if (row_i % 2 == 0) else fill_time
                for cell_item in [c_num, c_date, c_day, c_slot, c_time, c_subj, c_teach, c_room, c_hrs]:
                    cell_item.font = font_content
                    cell_item.border = cell_border
                    cell_item.fill = fill_curr

                c_num.alignment = align_center
                c_date.alignment = align_center
                c_day.alignment = align_center
                c_slot.alignment = align_center
                c_time.alignment = align_center
                c_subj.alignment = Alignment(horizontal="left", vertical="center")
                c_teach.alignment = Alignment(horizontal="left", vertical="center")
                c_room.alignment = align_center
                c_hrs.alignment = align_center

            # Summary Row
            sum_row = 5 + len(group_slots)
            ws_annual.row_dimensions[sum_row].height = 24
            ws_annual.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=8)
            sum_c1 = ws_annual.cell(row=sum_row, column=1, value="ИТОГО ПОЛНАЯ НАГРУЗКА ГОДА (ЧАСОВ):")
            sum_c1.font = Font(name="Arial", size=10, bold=True, color="0F172A")
            sum_c1.alignment = Alignment(horizontal="right", vertical="center")
            sum_c2 = ws_annual.cell(row=sum_row, column=9, value=total_slot_hours)
            sum_c2.font = Font(name="Arial", size=10, bold=True, color="0F172A")
            sum_c2.alignment = align_center
            fill_sum = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")
            for c_idx in range(1, 10):
                cell_s = ws_annual.cell(row=sum_row, column=c_idx)
                cell_s.border = cell_border
                cell_s.fill = fill_sum

    # 5. Save to memory buffer
    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer
